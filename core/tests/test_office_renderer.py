from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

_OFFICE_RUNTIME_ROOT = (
    Path(__file__).resolve().parents[1] / "skills" / "office-renderer" / "scripts"
)
sys.path.insert(0, str(_OFFICE_RUNTIME_ROOT))

from lamtools_office_renderer import check_office, render_office, sha256_file, validate_office


def test_office_readiness_check_converts_three_minimal_formats(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from lamtools_office_renderer.models import BackendInfo

    converted: list[str] = []

    class FakeBackend:
        info = BackendInfo("fake_office", "1.0")

        def render_pdf(self, source: Path, destination: Path, *, timeout: float) -> None:
            assert timeout == 7.0
            converted.append(source.suffix)
            destination.write_bytes(b"%PDF-test")

    monkeypatch.setattr(
        "lamtools_office_renderer.service.select_backend",
        lambda source, requested, discovery_timeout: FakeBackend(),
    )

    report = check_office(timeout=7.0)

    assert report["ok"] is True
    assert report["message"] == "Office 应用已就绪 · Test Passed 3/3"
    assert report["passed"] == report["total"] == 3
    assert converted == [".docx", ".xlsx", ".pptx"]


def _manifest(source: Path, datasets: dict, bindings: list[dict], **extra: object) -> dict:
    return {
        "schema_version": 1,
        "source": {"path": str(source), "sha256": sha256_file(source)},
        "datasets": datasets,
        "bindings": bindings,
        "tolerance": {"absolute": 0.0, "relative": 0.0},
        **extra,
    }


def test_docx_text_table_validates_and_hash_mismatch_fails_structure(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    source = tmp_path / "text.docx"
    document = docx.Document()
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, (("Name", "Note"), ("A", "alpha"), ("B", "beta"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(source)
    manifest = _manifest(
        source,
        {"notes": {"columns": [{"name": "Name", "type": "string"}, {"name": "Note", "type": "string"}], "rows": [["A", "alpha"], ["B", "beta"]]}},
        [{"id": "notes-table", "type": "docx_table", "dataset": "notes", "target": {"table": 1}}],
    )

    report = validate_office(manifest)
    assert report.ok
    assert report.as_dict()["statuses"] == {"data": "passed", "structure": "passed", "visual": "not_run"}

    manifest["source"]["sha256"] = "0" * 64
    report = validate_office(manifest)
    assert report.exit_class == "structure_invalid"
    assert [issue.code for issue in report.issues] == ["source_hash_mismatch"]


def test_manifest_path_resolves_relative_source_and_report_is_json_serializable(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    source = tmp_path / "relative.docx"
    docx.Document().save(source)
    manifest = _manifest(source, {}, [])
    manifest["source"]["path"] = source.name
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    report = validate_office(manifest_path)
    assert report.ok
    assert report.input_path == str(source.resolve())
    json.dumps(report.as_dict(), sort_keys=True)


def test_tolerance_is_explicit_contract(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    source = tmp_path / "tolerance.docx"
    docx.Document().save(source)
    manifest = _manifest(source, {}, [])
    del manifest["tolerance"]
    # With no bindings there is nothing numeric to compare, so the top-level
    # tolerance is optional. Every binding, however, must resolve one.
    assert validate_office(manifest).ok
    document = docx.Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "A"
    document.save(source)
    manifest = {
        "schema_version": 1,
        "source": {"path": str(source), "sha256": sha256_file(source)},
        "datasets": {"d": {"columns": ["A"], "rows": []}},
        "bindings": [{"id": "b", "type": "docx_table", "dataset": "d", "target": {"table": 1}}],
    }
    report = validate_office(manifest)
    assert report.exit_class == "structure_invalid"
    assert report.issues[0].code == "tolerance_invalid"


def test_unhashable_binding_fields_are_structured_and_cli_exits_two(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    from lamtools_core.cli import OFFICE_EXIT_MANIFEST_MISMATCH, main

    source = tmp_path / "invalid-binding.docx"
    pytest.importorskip("docx").Document().save(source)
    manifest = _manifest(
        source,
        {"d": {"columns": ["A"], "rows": []}},
        [{"id": "bad", "type": ["docx_table"], "dataset": ["d"], "target": {"table": 1}}],
    )
    report = validate_office(manifest)
    assert report.exit_class == "structure_invalid"
    assert {issue.code for issue in report.issues} >= {"binding_type_unsupported", "binding_dataset_missing"}

    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    assert main(["office", "validate", str(source), str(manifest_path)]) == OFFICE_EXIT_MANIFEST_MISMATCH
    payload = json.loads(capsys.readouterr().out)
    assert payload["exit_class"] == "structure_invalid"


def test_table_mismatch_is_cell_scoped_and_render_fails_before_backend(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    docx = pytest.importorskip("docx")
    source = tmp_path / "mismatch.docx"
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ("Key", "Value")):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ("x", "wrong")):
        cell.text = value
    document.save(source)
    manifest = _manifest(source, {"d": {"columns": ["Key", "Value"], "rows": [["x", "right"]]}}, [{"id": "b", "type": "docx_table", "dataset": "d", "target": {"table": 1}}])

    def forbidden(*args: object, **kwargs: object) -> object:
        raise AssertionError("backend discovery must not run")

    monkeypatch.setattr("lamtools_office_renderer.service.select_backend", forbidden)
    report = render_office(manifest, tmp_path / "out")
    assert report.exit_class == "data_mismatch"
    issue = next(issue for issue in report.issues if issue.code == "table_cell_mismatch")
    assert issue.binding == "b"
    assert issue.target == {"table": 1}
    assert issue.location == {"cell": "B2", "row": 2, "column": 2}
    assert (issue.expected, issue.actual) == ("right", "wrong")


def test_unbound_table_requires_explicit_non_data_scope(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    source = tmp_path / "decorative.docx"
    document = docx.Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "layout"
    document.save(source)
    manifest = _manifest(source, {}, [])
    report = validate_office(manifest)
    assert report.exit_class == "structure_invalid"
    assert report.issues[0].code == "unbound_data_element"

    manifest["non_data_targets"] = [{"type": "docx_table", "target": {"table": 1}}]
    assert validate_office(manifest).ok


def test_non_data_target_must_match_complete_native_target(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    source = tmp_path / "non-data.pptx"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.table.cell(0, 0).text = "layout"
    presentation.save(source)
    manifest = _manifest(
        source,
        {},
        [],
        non_data_targets=[{"type": "pptx_table", "target": {"slide": 1}}],
    )
    report = validate_office(manifest)
    assert report.exit_class == "structure_invalid"
    assert any(issue.code == "unbound_data_element" for issue in report.issues)

    manifest["non_data_targets"] = [
        {
            "type": "pptx_table",
            "target": {"slide": 1, "shape_id": shape.shape_id, "shape_name": shape.name},
        }
    ]
    assert validate_office(manifest).ok


def test_pptx_chart_and_table_bindings_validate_numeric_and_text_data(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    source = tmp_path / "deck.pptx"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(0.5), Inches(0.5), Inches(3), Inches(1))
    for cell, value in zip((cell for row in table_shape.table.rows for cell in row.cells), ("Label", "Comment", "A", "text")):
        cell.text = value
    chart_data = ChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Sales", (10.0, 20.0))
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2), Inches(5), Inches(3), chart_data)
    presentation.save(source)
    manifest = _manifest(
        source,
        {
            "labels": {"columns": ["Label", "Comment"], "rows": [["A", "text"]]},
            "sales": {"columns": [{"name": "Quarter", "type": "string"}, {"name": "Sales", "type": "number", "unit": "USD"}], "rows": [["Q1", 10], ["Q2", 20]]},
        },
        [
            {"id": "labels", "type": "pptx_table", "dataset": "labels", "target": {"slide": 1, "shape_id": table_shape.shape_id}},
            {"id": "sales", "type": "pptx_chart", "dataset": "sales", "target": {"slide": 1, "shape_id": chart_shape.shape_id}, "category": "Quarter", "series": ["Sales"]},
        ],
    )
    report = validate_office(manifest)
    assert report.ok, report.as_dict()


def test_xlsx_range_and_chart_bindings_validate(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.table import Table

    source = tmp_path / "book.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in (("Quarter", "Sales"), ("Q1", 10), ("Q2", 20)):
        sheet.append(row)
    sheet.add_table(Table(displayName="SalesTable", ref="A1:B3"))
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    sheet.add_chart(chart, "D2")
    workbook.save(source)
    manifest = _manifest(
        source,
        {"sales": {"columns": ["Quarter", "Sales"], "rows": [["Q1", 10], ["Q2", 20]]}},
        [
            {"id": "range", "type": "xlsx_range", "dataset": "sales", "target": {"sheet": "Data", "range": "A1:B3"}},
            {"id": "chart", "type": "xlsx_chart", "dataset": "sales", "target": {"sheet": "Data", "chart": 1}, "series": ["Sales"]},
        ],
    )
    report = validate_office(manifest)
    assert report.ok, report.as_dict()


def test_pdf_passthrough_renders_pngs_and_diagnostics(tmp_path: Path) -> None:
    fitz = pytest.importorskip("fitz")
    source = tmp_path / "source.pdf"
    document = fitz.open()
    page = document.new_page(width=300, height=200)
    page.insert_text((30, 50), "Office renderer")
    document.save(source)
    document.close()
    manifest = _manifest(source, {}, [])

    report = render_office(manifest, tmp_path / "rendered")
    assert report.ok, report.as_dict()
    assert report.backend and report.backend.name == "pdf_passthrough"
    assert Path(report.pdf_path or "").is_file()
    assert len(report.png_paths) == 1 and Path(report.png_paths[0]).is_file()
    assert report.text_diagnostics[0]["text"] == "Office renderer"


def test_backend_unavailable_is_explicit_but_validate_only_succeeds(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    docx = pytest.importorskip("docx")
    from lamtools_office_renderer import OfficeBackendUnavailable

    source = tmp_path / "empty.docx"
    docx.Document().save(source)
    manifest = _manifest(source, {}, [])
    assert validate_office(manifest).ok

    def unavailable(*args: object, **kwargs: object) -> object:
        raise OfficeBackendUnavailable("none installed")

    monkeypatch.setattr("lamtools_office_renderer.service.select_backend", unavailable)
    report = render_office(manifest, tmp_path / "out")
    assert report.exit_class == "visual_unavailable"
    assert report.as_dict()["statuses"] == {"data": "passed", "structure": "passed", "visual": "unavailable"}
    assert report.issues[0].code == "visual_backend_unavailable"


def test_auto_backend_falls_back_after_com_conversion_failure(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pymupdf = pytest.importorskip("pymupdf")
    from lamtools_office_renderer import BackendInfo, OfficeRendererError

    source = tmp_path / "empty.docx"
    pytest.importorskip("docx").Document().save(source)
    manifest = _manifest(source, {}, [])
    calls: list[str] = []

    class FakeBackend:
        def __init__(self, name: str) -> None:
            self.info = BackendInfo(name, "test")

        def render_pdf(self, source_path: Path, destination: Path, *, timeout: float) -> None:
            calls.append(self.info.name)
            if self.info.name == "microsoft_office_com":
                raise OfficeRendererError("COM export failed")
            document = pymupdf.open()
            document.new_page()
            destination.parent.mkdir(parents=True, exist_ok=True)
            document.save(destination)
            document.close()

    def choose(source_path: Path, *, requested: str, discovery_timeout: float) -> FakeBackend:
        del source_path, discovery_timeout
        return FakeBackend("microsoft_office_com" if requested == "auto" else "libreoffice")

    monkeypatch.setattr("lamtools_office_renderer.service.select_backend", choose)
    report = render_office(manifest, tmp_path / "out", backend="auto")
    assert report.ok, report.as_dict()
    assert calls == ["microsoft_office_com", "libreoffice"]
    assert report.backend and report.backend.name == "libreoffice"


def test_auto_backend_timeout_does_not_start_fallback(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from lamtools_office_renderer import BackendInfo, OfficeRenderTimeout

    source = tmp_path / "timeout.docx"
    pytest.importorskip("docx").Document().save(source)
    manifest = _manifest(source, {}, [])
    discoveries: list[str] = []

    class TimeoutBackend:
        info = BackendInfo("microsoft_office_com", "test")

        def render_pdf(self, source_path: Path, destination: Path, *, timeout: float) -> None:
            del source_path, destination, timeout
            raise OfficeRenderTimeout("timed out")

    def choose(source_path: Path, *, requested: str, discovery_timeout: float) -> TimeoutBackend:
        del source_path, discovery_timeout
        discoveries.append(requested)
        return TimeoutBackend()

    monkeypatch.setattr("lamtools_office_renderer.service.select_backend", choose)
    report = render_office(manifest, tmp_path / "out", backend="auto")
    assert report.exit_class == "visual_issue"
    assert report.issues[0].code == "render_timeout"
    assert discoveries == ["auto"]


def test_windows_libreoffice_discovery_prefers_console_launcher(monkeypatch: pytest.MonkeyPatch) -> None:
    import lamtools_office_renderer.backends as backend_module

    launched: list[str] = []
    console = Path(r"C:\Program Files\LibreOffice\program\soffice.com")
    gui = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")

    monkeypatch.setattr(backend_module.os, "name", "nt")
    monkeypatch.setattr(backend_module.shutil, "which", lambda name: None)
    monkeypatch.setattr(backend_module.Path, "is_file", lambda path: path in {console, gui})

    def run(command: list[str], *, timeout: float) -> object:
        del timeout
        launched.append(command[0])
        return backend_module.subprocess.CompletedProcess(command, 0, "LibreOffice 25.2.5.2\n", "")

    monkeypatch.setattr(backend_module, "_run_process", run)
    backend = backend_module.LibreOfficeBackend.discover()
    assert backend is not None
    assert backend.executable == str(console)
    assert launched == [str(console)]
    assert backend.info.version == "LibreOffice 25.2.5.2"


def test_windows_path_libreoffice_exe_yields_to_sibling_console_launcher(monkeypatch: pytest.MonkeyPatch) -> None:
    import lamtools_office_renderer.backends as backend_module

    gui = Path(r"D:\LibreOffice\program\soffice.exe")
    console = gui.with_name("soffice.com")
    launched: list[str] = []
    monkeypatch.setattr(backend_module.os, "name", "nt")
    monkeypatch.setattr(backend_module.shutil, "which", lambda name: str(gui) if name == "soffice" else None)
    monkeypatch.setattr(backend_module.Path, "is_file", lambda path: path in {console, gui})

    def run(command: list[str], *, timeout: float) -> object:
        del timeout
        launched.append(command[0])
        return backend_module.subprocess.CompletedProcess(command, 0, "LibreOffice PATH\n", "")

    monkeypatch.setattr(backend_module, "_run_process", run)
    backend = backend_module.LibreOfficeBackend.discover()
    assert backend is not None and backend.executable == str(console)
    assert launched == [str(console)]


def test_explicit_libreoffice_backend_smoke(tmp_path: Path) -> None:
    from lamtools_office_renderer.backends import LibreOfficeBackend

    if LibreOfficeBackend.discover() is None:
        pytest.skip("LibreOffice is not installed on this host")
    source = tmp_path / "libreoffice.docx"
    document = pytest.importorskip("docx").Document()
    document.add_paragraph("LibreOffice smoke")
    document.save(source)
    manifest = _manifest(source, {}, [])

    report = render_office(manifest, tmp_path / "out", backend="libreoffice", timeout=30)
    assert report.ok, report.as_dict()
    assert report.backend and report.backend.name == "libreoffice"
    assert report.pdf_path and Path(report.pdf_path).is_file()
    assert report.png_paths and all(Path(path).is_file() for path in report.png_paths)


def test_pptx_geometry_overlap_issue_is_complete_and_allowlist_is_pair_scoped(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    from lamtools_office_renderer.visual import inspect_pptx_geometry

    source = tmp_path / "overlap.pptx"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    left = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    left.text = "left text"
    right = slide.shapes.add_textbox(Inches(2), Inches(1), Inches(3), Inches(1))
    right.text = "right text"
    presentation.save(source)

    issues = inspect_pptx_geometry(source, {}, tolerance=0.5)
    assert len(issues) == 1
    issue = issues[0]
    assert issue.page == 1 and issue.kind == "pptx_geometry"
    assert issue.elements == (f"pptx:shape:{left.shape_id}", f"pptx:shape:{right.shape_id}")
    assert issue.texts == ("left text", "right text")
    assert len(issue.bboxes) == 2 and issue.intersection is not None

    manifest = {"allow_overlaps": [{"page": 1, "kind": "pptx_geometry", "elements": list(issue.elements)}]}
    assert inspect_pptx_geometry(source, manifest, tolerance=0.5) == []


def test_pdf_span_overlap_detects_text_coalesced_into_one_block(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    pymupdf = pytest.importorskip("pymupdf")
    from lamtools_office_renderer.visual import inspect_pdf

    source = tmp_path / "span-overlap.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((50, 50), "first")
    page.insert_text((50, 50), "second")
    document.save(source)
    document.close()

    _, diagnostics, issues = inspect_pdf(source, tmp_path / "pages", source_path=source, manifest={})
    overlap = next(issue for issue in issues if issue.kind == "pdf_text")
    assert overlap.page == 1
    assert overlap.elements == ("pdf:span:0:0:0", "pdf:span:0:1:0")
    assert overlap.texts == ("first", "second")
    assert len(overlap.bboxes) == 2 and overlap.intersection is not None
    assert [item["kind"] for item in diagnostics] == ["text", "text"]
    assert capsys.readouterr().out == ""


def test_pdf_table_region_detects_non_cell_text_conflict(tmp_path: Path) -> None:
    pymupdf = pytest.importorskip("pymupdf")
    from lamtools_office_renderer.visual import inspect_pdf

    source = tmp_path / "table-conflict.pdf"
    document = pymupdf.open()
    page = document.new_page(width=300, height=300)
    for x in (40, 140, 240):
        page.draw_line((x, 40), (x, 140))
    for y in (40, 90, 140):
        page.draw_line((40, y), (240, y))
    for point, text in (((50, 70), "A"), ((150, 70), "B"), ((50, 120), "C"), ((150, 120), "D")):
        page.insert_text(point, text)
    page.insert_text((100, 135), "footer crossing")
    document.save(source)
    document.close()

    _, diagnostics, issues = inspect_pdf(source, tmp_path / "pages", source_path=source, manifest={})
    assert any(item["kind"] == "table" for item in diagnostics)
    issue = next(issue for issue in issues if issue.kind == "pdf_table_text")
    assert issue.page == 1
    assert issue.elements[0] == "pdf:table:0"
    assert issue.texts[1] == "footer crossing"
    assert issue.intersection is not None
