from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

from .models import OfficeIssue


@dataclass(frozen=True, slots=True)
class SourceValidation:
    data_issues: tuple[OfficeIssue, ...]
    structure_issues: tuple[OfficeIssue, ...]

    @property
    def ok(self) -> bool:
        return not self.data_issues and not self.structure_issues


@dataclass(frozen=True, slots=True)
class NativeElement:
    kind: str
    target: dict[str, Any]
    payload: Any


def validate_source(manifest: Mapping[str, Any], source_path: Path) -> SourceValidation:
    suffix = source_path.suffix.casefold()
    try:
        elements = _inventory(source_path)
    except ImportError as exc:
        return SourceValidation((), (OfficeIssue("office_dependency_unavailable", str(exc)),))
    except Exception as exc:
        return SourceValidation((), (OfficeIssue("office_source_unreadable", str(exc), actual=str(source_path)),))

    data_issues: list[OfficeIssue] = []
    structure_issues: list[OfficeIssue] = []
    bindings = manifest.get("bindings", [])
    claimed: set[int] = set()
    for binding in bindings:
        if not isinstance(binding, Mapping):
            continue
        binding_id = str(binding.get("id", ""))
        kind = str(binding.get("type", ""))
        target = dict(binding.get("target", {}))
        expected_suffix = {"pptx": ".pptx", "docx": ".docx", "xlsx": ".xlsx"}.get(kind.split("_", 1)[0])
        if expected_suffix != suffix:
            structure_issues.append(
                OfficeIssue(
                    "binding_source_type_mismatch",
                    "binding type does not apply to source file",
                    binding=binding_id,
                    target=target,
                    expected=expected_suffix,
                    actual=suffix,
                )
            )
            continue
        if kind == "xlsx_range" and not any(item.kind == kind and _target_matches(target, item.target) for item in elements):
            try:
                materialized = inventory_xlsx_range(source_path, target)
            except Exception as exc:
                structure_issues.append(OfficeIssue("binding_target_unreadable", str(exc), binding=binding_id, target=target))
                continue
            if materialized is not None:
                elements.append(materialized)
        matches = [(index, item) for index, item in enumerate(elements) if item.kind == kind and _target_matches(target, item.target)]
        if not matches:
            structure_issues.append(
                OfficeIssue(
                    "binding_target_not_found",
                    "binding target was not found",
                    binding=binding_id,
                    target=target,
                    page=_page(target),
                    expected=target,
                    actual=None,
                )
            )
            continue
        if len(matches) > 1:
            structure_issues.append(
                OfficeIssue(
                    "binding_target_ambiguous",
                    "binding target identifies more than one element",
                    binding=binding_id,
                    target=target,
                    page=_page(target),
                    expected=1,
                    actual=len(matches),
                )
            )
            continue
        index, element = matches[0]
        if index in claimed:
            structure_issues.append(OfficeIssue("binding_target_duplicate", "multiple bindings claim the same target", binding=binding_id, target=element.target, page=_page(element.target)))
            continue
        claimed.add(index)
        dataset = manifest["datasets"][binding["dataset"]]
        tolerance = binding.get("tolerance", manifest.get("tolerance", {}))
        absolute = float(tolerance.get("absolute", 0.0))
        relative = float(tolerance.get("relative", 0.0))
        try:
            if kind.endswith("_table") or kind == "xlsx_range":
                expected = _dataset_matrix(dataset, include_header=bool(binding.get("include_header", True)))
                data_issues.extend(_compare_matrix(binding_id, element.target, expected, element.payload, absolute, relative))
            else:
                expected = _dataset_chart(dataset, binding)
                data_issues.extend(_compare_chart(binding_id, element.target, expected, element.payload, absolute, relative))
        except (KeyError, IndexError, TypeError, ValueError) as exc:
            structure_issues.append(OfficeIssue("binding_mapping_invalid", str(exc), binding=binding_id, target=element.target, page=_page(element.target)))

    non_data = manifest.get("non_data_targets", [])
    for index, element in enumerate(elements):
        if index in claimed:
            continue
        if any(_non_data_matches(scope, element) for scope in non_data):
            continue
        structure_issues.append(
            OfficeIssue(
                "unbound_data_element",
                "every table and chart must be bound or explicitly declared non-data",
                target=element.target,
                page=_page(element.target),
                expected="binding or non_data_targets entry",
                actual=element.kind,
            )
        )

    return SourceValidation(tuple(data_issues), tuple(structure_issues))


def _inventory(path: Path) -> list[NativeElement]:
    suffix = path.suffix.casefold()
    if suffix == ".pptx":
        return _inventory_pptx(path)
    if suffix == ".docx":
        return _inventory_docx(path)
    if suffix == ".xlsx":
        return _inventory_xlsx(path)
    if suffix == ".pdf":
        return []
    raise ValueError(f"unsupported Office source type: {suffix or '<none>'}")


def _inventory_pptx(path: Path) -> list[NativeElement]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError("PPTX validation requires python-pptx") from exc
    presentation = Presentation(path)
    found: list[NativeElement] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            base = {"slide": slide_number, "shape_id": int(shape.shape_id), "shape_name": shape.name}
            if getattr(shape, "has_table", False):
                matrix = [[cell.text for cell in row.cells] for row in shape.table.rows]
                found.append(NativeElement("pptx_table", base, matrix))
            if getattr(shape, "has_chart", False):
                found.append(NativeElement("pptx_chart", base, _pptx_chart_payload(shape.chart)))
    return found


def _pptx_chart_payload(chart: Any) -> dict[str, Any]:
    categories: list[Any] = []
    series: list[dict[str, Any]] = []
    if chart.plots:
        categories = [_category_value(category) for category in chart.plots[0].categories]
    for item in chart.series:
        series.append({"name": item.name, "values": list(item.values)})
    return {"categories": categories, "series": series}


def _category_value(category: Any) -> Any:
    label = getattr(category, "label", None)
    return label if label is not None else str(category)


def _inventory_docx(path: Path) -> list[NativeElement]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - declared dependency
        raise ImportError("DOCX validation requires python-docx") from exc
    document = Document(path)
    return [
        NativeElement(
            "docx_table",
            {"table": index},
            [[cell.text for cell in row.cells] for row in table.rows],
        )
        for index, table in enumerate(document.tables, 1)
    ]


def _inventory_xlsx(path: Path) -> list[NativeElement]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError("XLSX validation requires openpyxl") from exc
    workbook = load_workbook(path, data_only=True, read_only=False)
    found: list[NativeElement] = []
    try:
        for sheet in workbook.worksheets:
            for table in sheet.tables.values():
                min_col, min_row, max_col, max_row = range_boundaries(table.ref)
                payload = [
                    [sheet.cell(row=row, column=column).value for column in range(min_col, max_col + 1)]
                    for row in range(min_row, max_row + 1)
                ]
                canonical = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(max_row, max_col).coordinate}"
                found.append(
                    NativeElement(
                        "xlsx_range",
                        {"sheet": sheet.title, "range": canonical, "table_name": table.name},
                        payload,
                    )
                )
            for chart_index, chart in enumerate(sheet._charts, 1):  # openpyxl exposes no public chart collection
                title = _xlsx_chart_title(chart)
                target = {"sheet": sheet.title, "chart": chart_index, "chart_name": title}
                found.append(NativeElement("xlsx_chart", target, _xlsx_chart_payload(workbook, sheet, chart)))
    finally:
        workbook.close()
    return found


def inventory_xlsx_range(path: Path, target: Mapping[str, Any]) -> NativeElement | None:
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    sheet_name = target.get("sheet")
    cell_range = target.get("range")
    if not isinstance(sheet_name, str) or not isinstance(cell_range, str):
        return None
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            return None
        sheet = workbook[sheet_name]
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        payload = [
            [sheet.cell(row=row, column=column).value for column in range(min_col, max_col + 1)]
            for row in range(min_row, max_row + 1)
        ]
        canonical = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(max_row, max_col).coordinate}"
        return NativeElement("xlsx_range", {"sheet": sheet_name, "range": canonical}, payload)
    finally:
        workbook.close()


def _xlsx_chart_payload(workbook: Any, default_sheet: Any, chart: Any) -> dict[str, Any]:
    categories: list[Any] = []
    series: list[dict[str, Any]] = []
    for position, item in enumerate(chart.ser):
        category_ref = _chart_reference(getattr(item, "cat", None) or getattr(item, "xVal", None))
        value_ref = _chart_reference(getattr(item, "val", None) or getattr(item, "yVal", None))
        if position == 0 and category_ref:
            categories = _read_formula_values(workbook, default_sheet, category_ref)
        name = _series_name(workbook, default_sheet, item, position + 1)
        values = _read_formula_values(workbook, default_sheet, value_ref) if value_ref else []
        series.append({"name": name, "values": values})
    return {"categories": categories, "series": series}


def _chart_reference(holder: Any) -> str | None:
    if holder is None:
        return None
    for name in ("numRef", "strRef"):
        ref = getattr(holder, name, None)
        formula = getattr(ref, "f", None) if ref is not None else None
        if formula:
            return formula
    return None


_FORMULA_RE = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$")


def _read_formula_values(workbook: Any, default_sheet: Any, formula: str) -> list[Any]:
    from openpyxl.utils.cell import column_index_from_string

    match = _FORMULA_RE.match(formula)
    if not match:
        return []
    sheet_name = (match.group(1) or match.group(2)).replace("''", "'")
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else default_sheet
    c1, r1, c2, r2 = column_index_from_string(match.group(3)), int(match.group(4)), column_index_from_string(match.group(5)), int(match.group(6))
    return [sheet.cell(row=row, column=column).value for row in range(r1, r2 + 1) for column in range(c1, c2 + 1)]


def _series_name(workbook: Any, default_sheet: Any, item: Any, position: int) -> Any:
    title = getattr(item, "tx", None)
    if title is not None:
        literal = getattr(title, "v", None)
        if literal is not None:
            return literal
        ref = getattr(getattr(title, "strRef", None), "f", None)
        if ref:
            values = _read_single_ref(workbook, default_sheet, ref)
            if values is not None:
                return values
    return f"Series {position}"


def _read_single_ref(workbook: Any, default_sheet: Any, formula: str) -> Any:
    from openpyxl.utils.cell import range_boundaries

    if "!" not in formula:
        return None
    raw_sheet, address = formula.rsplit("!", 1)
    sheet_name = raw_sheet.strip("'").replace("''", "'")
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else default_sheet
    min_col, min_row, _, _ = range_boundaries(address.replace("$", ""))
    return sheet.cell(min_row, min_col).value


def _xlsx_chart_title(chart: Any) -> str | None:
    title = getattr(chart, "title", None)
    try:
        return title.tx.rich.p[0].r[0].t if title and title.tx and title.tx.rich else None
    except (AttributeError, IndexError):
        return None


def _dataset_matrix(dataset: Mapping[str, Any], *, include_header: bool) -> list[list[Any]]:
    columns = [column if isinstance(column, str) else column["name"] for column in dataset["columns"]]
    rows = [list(row) for row in dataset["rows"]]
    return [columns, *rows] if include_header else rows


def _dataset_chart(dataset: Mapping[str, Any], binding: Mapping[str, Any]) -> dict[str, Any]:
    columns = [column if isinstance(column, str) else column["name"] for column in dataset["columns"]]
    category_key = binding.get("category", columns[0])
    category_index = _column_index(category_key, columns)
    requested_series = binding.get("series", columns[1:])
    if not isinstance(requested_series, list):
        requested_series = [requested_series]
    result = {"categories": [row[category_index] for row in dataset["rows"]], "series": []}
    for spec in requested_series:
        if isinstance(spec, Mapping):
            column_key = spec.get("column")
            name = spec.get("name", column_key)
        else:
            column_key = spec
            name = spec
        index = _column_index(column_key, columns)
        result["series"].append({"name": name, "values": [row[index] for row in dataset["rows"]]})
    return result


def _column_index(value: Any, columns: Sequence[str]) -> int:
    if isinstance(value, int):
        return value
    if value in columns:
        return columns.index(value)
    raise ValueError(f"unknown dataset column: {value!r}")


def _compare_matrix(binding: str, target: dict[str, Any], expected: list[list[Any]], actual: list[list[Any]], absolute: float, relative: float) -> list[OfficeIssue]:
    issues: list[OfficeIssue] = []
    if len(expected) != len(actual):
        issues.append(_mismatch("table_row_count_mismatch", binding, target, {"dimension": "rows"}, len(expected), len(actual)))
    expected_width = max((len(row) for row in expected), default=0)
    actual_width = max((len(row) for row in actual), default=0)
    if expected_width != actual_width:
        issues.append(_mismatch("table_column_count_mismatch", binding, target, {"dimension": "columns"}, expected_width, actual_width))
    for row_index in range(min(len(expected), len(actual))):
        for column_index in range(min(len(expected[row_index]), len(actual[row_index]))):
            if not _equal(expected[row_index][column_index], actual[row_index][column_index], absolute, relative):
                issues.append(_mismatch("table_cell_mismatch", binding, target, {"cell": _cell_label(row_index, column_index), "row": row_index + 1, "column": column_index + 1}, expected[row_index][column_index], actual[row_index][column_index]))
    return issues


def _compare_chart(binding: str, target: dict[str, Any], expected: dict[str, Any], actual: dict[str, Any], absolute: float, relative: float) -> list[OfficeIssue]:
    issues: list[OfficeIssue] = []
    expected_categories = expected["categories"]
    actual_categories = actual["categories"]
    if len(expected_categories) != len(actual_categories):
        issues.append(_mismatch("chart_category_count_mismatch", binding, target, {"category": "count"}, len(expected_categories), len(actual_categories)))
    for index, (left, right) in enumerate(zip(expected_categories, actual_categories)):
        if not _equal(left, right, absolute, relative):
            issues.append(_mismatch("chart_category_mismatch", binding, target, {"category": index + 1}, left, right))
    expected_series = expected["series"]
    actual_series = actual["series"]
    if len(expected_series) != len(actual_series):
        issues.append(_mismatch("chart_series_count_mismatch", binding, target, {"series": "count"}, len(expected_series), len(actual_series)))
    for series_index, (left, right) in enumerate(zip(expected_series, actual_series)):
        if not _equal(left["name"], right["name"], absolute, relative):
            issues.append(_mismatch("chart_series_name_mismatch", binding, target, {"series": series_index + 1, "field": "name"}, left["name"], right["name"]))
        if len(left["values"]) != len(right["values"]):
            issues.append(_mismatch("chart_value_count_mismatch", binding, target, {"series": series_index + 1, "field": "values"}, len(left["values"]), len(right["values"])))
        for category_index, (expected_value, actual_value) in enumerate(zip(left["values"], right["values"])):
            if not _equal(expected_value, actual_value, absolute, relative):
                issues.append(_mismatch("chart_value_mismatch", binding, target, {"series": series_index + 1, "category": category_index + 1}, expected_value, actual_value))
    return issues


def _equal(expected: Any, actual: Any, absolute: float, relative: float) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool) and isinstance(actual, (int, float)) and not isinstance(actual, bool):
        return math.isclose(float(expected), float(actual), abs_tol=absolute, rel_tol=relative)
    return expected == actual or str(expected) == str(actual)


def _mismatch(code: str, binding: str, target: dict[str, Any], location: dict[str, Any], expected: Any, actual: Any) -> OfficeIssue:
    return OfficeIssue(code, "rendered source data differs from canonical dataset", binding=binding, target=target, page=_page(target), location=location, expected=expected, actual=actual)


def _target_matches(expected: Mapping[str, Any], actual: Mapping[str, Any]) -> bool:
    return all(actual.get(key) == value for key, value in expected.items())


def _non_data_matches(scope: Any, element: NativeElement) -> bool:
    if not isinstance(scope, Mapping):
        return False
    kind = scope.get("type")
    target = scope.get("target")
    return kind == element.kind and isinstance(target, Mapping) and dict(target) == element.target


def _page(target: Mapping[str, Any]) -> int | None:
    value = target.get("slide") or target.get("page")
    return int(value) if isinstance(value, int) else None


def _cell_label(row: int, column: int) -> str:
    name = ""
    number = column + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        name = chr(65 + remainder) + name
    return f"{name}{row + 1}"
