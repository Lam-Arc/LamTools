from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any, Mapping

from .backends import select_backend
from .manifest import ManifestValidation, sha256_file, validate_manifest_shape
from .models import (
    OfficeBackendUnavailable,
    OfficeIssue,
    OfficeRenderReport,
    OfficeRendererError,
    OfficeRenderTimeout,
    normalized_path,
)
from .validation import validate_source
from .visual import inspect_pdf


def check_office(*, backend: str = "auto", timeout: float = 15.0) -> dict[str, Any]:
    """Run one bounded readiness check across Word, Excel, and PowerPoint formats."""

    if backend not in {"auto", "microsoft", "libreoffice"}:
        raise ValueError("backend must be auto, microsoft, or libreoffice")
    tests: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix="lamtools-office-check-") as temp_name:
        root = Path(temp_name)
        sources = _create_check_sources(root)
        for label, source in sources:
            destination = root / f"{source.stem}.pdf"
            try:
                selected = select_backend(
                    source,
                    requested=backend,
                    discovery_timeout=min(5.0, timeout),
                )
                try:
                    selected.render_pdf(source, destination, timeout=timeout)
                except (OfficeRendererError, OfficeRenderTimeout):
                    if backend != "auto" or selected.info.name != "microsoft_office_com":
                        raise
                    selected = select_backend(
                        source,
                        requested="libreoffice",
                        discovery_timeout=min(5.0, timeout),
                    )
                    selected.render_pdf(source, destination, timeout=timeout)
                if not destination.is_file() or destination.stat().st_size < 5:
                    raise OfficeRendererError("conversion did not produce a PDF")
                tests.append(
                    {
                        "name": label,
                        "status": "passed",
                        "backend": selected.info.as_dict(),
                    }
                )
            except (OfficeBackendUnavailable, OfficeRendererError, OfficeRenderTimeout) as exc:
                tests.append(
                    {
                        "name": label,
                        "status": "failed",
                        "error": {"type": type(exc).__name__, "message": str(exc)},
                    }
                )
    passed = sum(item["status"] == "passed" for item in tests)
    total = len(tests)
    ok = passed == total
    readiness = "已就绪" if ok else "未就绪"
    return {
        "ok": ok,
        "exit_class": "success" if ok else "render_unavailable",
        "message": f"Office 应用{readiness} · Test Passed {passed}/{total}",
        "passed": passed,
        "total": total,
        "tests": tests,
    }


def _create_check_sources(root: Path) -> list[tuple[str, Path]]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    docx_path = root / "word-check.docx"
    document = Document()
    document.add_paragraph("LamTools Office readiness check")
    document.save(docx_path)

    xlsx_path = root / "excel-check.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "LamTools Office readiness check"
    sheet["A2"] = 1
    sheet["A3"] = "=A2+1"
    workbook.save(xlsx_path)

    pptx_path = root / "powerpoint-check.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    box.text = "LamTools Office readiness check"
    presentation.save(pptx_path)

    return [("Word", docx_path), ("Excel", xlsx_path), ("PowerPoint", pptx_path)]


def validate_office(
    manifest: Mapping[str, Any] | str | Path,
    *,
    source: str | Path | None = None,
) -> OfficeRenderReport:
    checked = validate_manifest_shape(manifest, source_override=source)
    return _report_from_validation(checked, source=source)


def _report_from_validation(checked: ManifestValidation, *, source: str | Path | None) -> OfficeRenderReport:
    input_path = normalized_path(checked.source_path) if checked.source_path else str(source or "")
    if checked.issues:
        return OfficeRenderReport(
            schema_version=1,
            input_path=input_path,
            input_sha256=checked.source_hash or "",
            backend=None,
            pdf_path=None,
            png_paths=[],
            data_status="not_run",
            structure_status="failed",
            visual_status="not_run",
            issues=list(checked.issues),
        )
    assert checked.source_path is not None
    source_check = validate_source(checked.manifest, checked.source_path)
    return OfficeRenderReport(
        schema_version=1,
        input_path=normalized_path(checked.source_path),
        input_sha256=checked.source_hash or "",
        backend=None,
        pdf_path=None,
        png_paths=[],
        data_status="failed" if source_check.data_issues else "passed",
        structure_status="failed" if source_check.structure_issues else "passed",
        visual_status="not_run",
        issues=[*source_check.structure_issues, *source_check.data_issues],
    )


def render_office(
    manifest: Mapping[str, Any] | str | Path,
    output_dir: str | Path,
    *,
    source: str | Path | None = None,
    backend: str = "auto",
    timeout: float = 120.0,
    discovery_timeout: float = 8.0,
    dpi: int = 144,
    overlap_tolerance_points: float = 0.5,
) -> OfficeRenderReport:
    checked = validate_manifest_shape(manifest, source_override=source)
    report = _report_from_validation(checked, source=source)
    # This gate is intentionally before backend discovery and all child process
    # work: a data/structure mismatch must fail closed without rendering.
    if report.data_status != "passed" or report.structure_status != "passed":
        return report
    if backend not in {"auto", "microsoft", "libreoffice"}:
        report.structure_status = "failed"
        report.issues.append(OfficeIssue("backend_invalid", "backend must be auto, microsoft, or libreoffice", expected=["auto", "microsoft", "libreoffice"], actual=backend))
        return report
    source_path = Path(report.input_path)
    destination_dir = Path(output_dir).expanduser().resolve()
    pdf_path = destination_dir / f"{source_path.stem}.pdf"
    pages_dir = destination_dir / f"{source_path.stem}-pages"
    current_hash = sha256_file(source_path)
    if current_hash != report.input_sha256:
        report.structure_status = "failed"
        report.visual_status = "not_run"
        report.issues.append(OfficeIssue("source_changed", "source changed after validation and before rendering", expected=report.input_sha256, actual=current_hash))
        return report
    try:
        selected_backend = select_backend(source_path, requested=backend, discovery_timeout=discovery_timeout)
    except OfficeBackendUnavailable as exc:
        report.visual_status = "unavailable"
        report.issues.append(OfficeIssue("visual_backend_unavailable", str(exc), actual=source_path.suffix.casefold()))
        return report
    report.backend = selected_backend.info
    try:
        selected_backend.render_pdf(source_path, pdf_path, timeout=timeout)
    except OfficeRenderTimeout as exc:
        pdf_path.unlink(missing_ok=True)
        report.visual_status = "failed"
        report.issues.append(OfficeIssue("render_timeout", str(exc), actual=timeout))
        return report
    except OfficeRendererError as exc:
        pdf_path.unlink(missing_ok=True)
        if backend == "auto" and selected_backend.info.name == "microsoft_office_com":
            try:
                fallback = select_backend(source_path, requested="libreoffice", discovery_timeout=discovery_timeout)
                fallback.render_pdf(source_path, pdf_path, timeout=timeout)
                report.backend = fallback.info
            except OfficeRenderTimeout as fallback_exc:
                pdf_path.unlink(missing_ok=True)
                report.visual_status = "failed"
                report.issues.append(OfficeIssue("render_timeout", str(fallback_exc), actual=timeout))
                return report
            except (OfficeBackendUnavailable, OfficeRendererError) as fallback_exc:
                pdf_path.unlink(missing_ok=True)
                report.visual_status = "failed"
                report.issues.append(OfficeIssue("render_failed", f"Microsoft Office failed: {exc}; LibreOffice fallback failed: {fallback_exc}"))
                return report
        else:
            pdf_path.unlink(missing_ok=True)
            report.visual_status = "failed"
            report.issues.append(OfficeIssue("render_failed", str(exc)))
            return report
    report.pdf_path = normalized_path(pdf_path)
    final_hash = sha256_file(source_path)
    if final_hash != report.input_sha256:
        pdf_path.unlink(missing_ok=True)
        report.structure_status = "failed"
        report.visual_status = "not_run"
        report.pdf_path = None
        report.issues.append(OfficeIssue("source_changed", "source changed while rendering", expected=report.input_sha256, actual=final_hash))
        return report
    try:
        pngs, diagnostics, visual_issues = inspect_pdf(
            pdf_path,
            pages_dir,
            source_path=source_path,
            manifest=checked.manifest,
            dpi=dpi,
            overlap_tolerance_points=overlap_tolerance_points,
        )
    except ImportError as exc:
        report.visual_status = "unavailable"
        report.issues.append(OfficeIssue("visual_dependency_unavailable", str(exc)))
        return report
    except Exception as exc:
        report.visual_status = "failed"
        report.issues.append(OfficeIssue("visual_inspection_failed", str(exc)))
        return report
    report.png_paths = [normalized_path(path) for path in pngs]
    report.text_diagnostics = diagnostics
    report.issues.extend(visual_issues)
    report.visual_status = "failed" if visual_issues else "passed"
    return report
