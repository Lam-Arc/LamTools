from __future__ import annotations

import contextlib
import io
from itertools import combinations
from pathlib import Path
from typing import Any, Mapping

from .models import OfficeIssue


def inspect_pdf(
    pdf_path: Path,
    output_dir: Path,
    *,
    source_path: Path,
    manifest: Mapping[str, Any],
    dpi: int = 144,
    overlap_tolerance_points: float = 0.5,
) -> tuple[list[Path], list[dict[str, Any]], list[OfficeIssue]]:
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError("visual inspection requires PyMuPDF") from exc
    output_dir.mkdir(parents=True, exist_ok=True)
    pngs: list[Path] = []
    diagnostics: list[dict[str, Any]] = []
    issues: list[OfficeIssue] = []
    page_texts: dict[int, str] = {}
    with pymupdf.open(pdf_path) as document:
        for page_index, page in enumerate(document, 1):
            pixmap = page.get_pixmap(dpi=dpi, alpha=False)
            png = output_dir / f"page-{page_index:04d}.png"
            pixmap.save(png)
            pngs.append(png)
            spans: list[tuple[str, str, tuple[float, float, float, float]]] = []
            page_texts[page_index] = " ".join(page.get_text("text").split())
            text_dict = page.get_text("dict")
            for block_index, block in enumerate(text_dict.get("blocks", [])):
                for line_index, line in enumerate(block.get("lines", [])):
                    for span_index, span in enumerate(line.get("spans", [])):
                        text = " ".join(str(span.get("text", "")).split())
                        if not text:
                            continue
                        bbox = _round_box(span["bbox"])
                        element_id = f"pdf:span:{block_index}:{line_index}:{span_index}"
                        diagnostics.append({"page": page_index, "element": element_id, "kind": "text", "text": text, "bbox": list(bbox)})
                        spans.append((element_id, text, bbox))
            for left, right in combinations(spans, 2):
                intersection = _intersection(left[2], right[2], overlap_tolerance_points)
                if intersection and not _allowed(manifest.get("allow_overlaps", []), page_index, "pdf_text", left[0], right[0]):
                    issues.append(_overlap_issue(page_index, "pdf_text", left, right, intersection))
            table_elements = _pdf_table_elements(page, page_index, diagnostics)
            for table, cell_boxes in table_elements:
                for span in spans:
                    intersection = _intersection(table[2], span[2], overlap_tolerance_points)
                    if not intersection or _span_belongs_to_table(span, table, cell_boxes):
                        continue
                    if not _allowed(manifest.get("allow_overlaps", []), page_index, "pdf_table_text", table[0], span[0]):
                        issues.append(_overlap_issue(page_index, "pdf_table_text", table, span, intersection))
            for (left, _), (right, _) in combinations(table_elements, 2):
                intersection = _intersection(left[2], right[2], overlap_tolerance_points)
                if intersection and not _allowed(manifest.get("allow_overlaps", []), page_index, "pdf_tables", left[0], right[0]):
                    issues.append(_overlap_issue(page_index, "pdf_tables", left, right, intersection))
    issues.extend(_rendered_table_text_issues(manifest, page_texts))
    if source_path.suffix.casefold() == ".pptx":
        issues.extend(inspect_pptx_geometry(source_path, manifest, tolerance=overlap_tolerance_points))
    issues.sort(key=lambda item: (item.page or 0, item.kind or "", item.elements))
    diagnostics.sort(key=lambda item: (item["page"], item["bbox"][1], item["bbox"][0], item["element"]))
    return pngs, diagnostics, issues


def _rendered_table_text_issues(manifest: Mapping[str, Any], page_texts: Mapping[int, str]) -> list[OfficeIssue]:
    issues: list[OfficeIssue] = []
    whole_document = " ".join(page_texts.values())
    datasets = manifest.get("datasets", {})
    for binding in manifest.get("bindings", []):
        if not isinstance(binding, Mapping) or binding.get("type") not in {"pptx_table", "docx_table", "xlsx_range"}:
            continue
        dataset = datasets.get(binding.get("dataset"), {})
        columns = [column if isinstance(column, str) else column.get("name") for column in dataset.get("columns", [])]
        values: list[tuple[str, Any]] = []
        if binding.get("include_header", True):
            values.extend((_cell_label(0, column), value) for column, value in enumerate(columns))
        row_offset = 1 if binding.get("include_header", True) else 0
        values.extend(
            (_cell_label(row + row_offset, column), value)
            for row, row_values in enumerate(dataset.get("rows", []))
            for column, value in enumerate(row_values)
        )
        target = dict(binding.get("target", {}))
        page_number = target.get("slide") if binding.get("type") == "pptx_table" else None
        haystack = page_texts.get(page_number, "") if isinstance(page_number, int) else whole_document
        for cell, value in values:
            needle = " ".join(str(value).split()) if value is not None else ""
            if needle and needle not in haystack:
                issues.append(
                    OfficeIssue(
                        "rendered_table_text_missing",
                        "validated table text was not found in the rendered PDF",
                        binding=str(binding.get("id", "")),
                        target=target,
                        page=page_number,
                        location={"cell": cell},
                        expected=needle,
                        actual=None,
                        kind="pdf_table_text",
                    )
                )
    return issues


def _pdf_table_elements(page: Any, page_number: int, diagnostics: list[dict[str, Any]]) -> list[tuple[tuple[str, str, tuple[float, float, float, float]], list[tuple[float, float, float, float]]]]:
    elements: list[tuple[tuple[str, str, tuple[float, float, float, float]], list[tuple[float, float, float, float]]]] = []
    try:
        # PyMuPDF currently prints an optional-layout-package advisory on first
        # use. A renderer API must keep stdout clean for its JSON CLI adapter.
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            tables = page.find_tables().tables
    except (AttributeError, RuntimeError, ValueError):
        return elements
    for index, table in enumerate(tables):
        extracted = table.extract()
        text = " | ".join(" ".join(str(value).split()) for row in extracted for value in row if value is not None and str(value).strip())
        element_id = f"pdf:table:{index}"
        bbox = _round_box(table.bbox)
        cells = [_round_box(cell) for cell in table.cells if cell is not None]
        diagnostics.append({"page": page_number, "element": element_id, "kind": "table", "text": text, "bbox": list(bbox)})
        elements.append(((element_id, text or "[table]", bbox), cells))
    return elements


def _span_belongs_to_table(
    span: tuple[str, str, tuple[float, float, float, float]],
    table: tuple[str, str, tuple[float, float, float, float]],
    cell_boxes: list[tuple[float, float, float, float]],
) -> bool:
    center = ((span[2][0] + span[2][2]) / 2, (span[2][1] + span[2][3]) / 2)
    in_cell = any(cell[0] <= center[0] <= cell[2] and cell[1] <= center[1] <= cell[3] for cell in cell_boxes)
    return in_cell and span[1] in table[1]


def inspect_pptx_geometry(source_path: Path, manifest: Mapping[str, Any], *, tolerance: float) -> list[OfficeIssue]:
    from pptx import Presentation

    presentation = Presentation(source_path)
    issues: list[OfficeIssue] = []
    for page_number, slide in enumerate(presentation.slides, 1):
        elements: list[tuple[str, str, tuple[float, float, float, float], str]] = []
        for shape in slide.shapes:
            kind, text = _shape_content(shape)
            if not text:
                continue
            element_id = f"pptx:shape:{shape.shape_id}"
            bbox = _round_box((shape.left / 12700.0, shape.top / 12700.0, (shape.left + shape.width) / 12700.0, (shape.top + shape.height) / 12700.0))
            elements.append((element_id, text, bbox, kind))
        for left, right in combinations(elements, 2):
            if left[3] != "text" and right[3] != "text":
                continue
            intersection = _intersection(left[2], right[2], tolerance)
            if intersection and not _allowed(manifest.get("allow_overlaps", []), page_number, "pptx_geometry", left[0], right[0]):
                issues.append(_overlap_issue(page_number, "pptx_geometry", left[:3], right[:3], intersection))
    return issues


def _shape_content(shape: Any) -> tuple[str, str]:
    if getattr(shape, "has_table", False):
        return "table", " | ".join(cell.text for row in shape.table.rows for cell in row.cells if cell.text.strip())
    if getattr(shape, "has_chart", False):
        return "chart", f"[chart:{shape.name}]"
    if getattr(shape, "has_text_frame", False):
        return "text", " ".join(shape.text.split())
    return "other", ""


def _intersection(left: tuple[float, float, float, float], right: tuple[float, float, float, float], tolerance: float) -> tuple[float, float, float, float] | None:
    box = (max(left[0], right[0]), max(left[1], right[1]), min(left[2], right[2]), min(left[3], right[3]))
    if box[2] - box[0] <= tolerance or box[3] - box[1] <= tolerance:
        return None
    return _round_box(box)


def _allowed(scopes: Any, page: int, kind: str, left: str, right: str) -> bool:
    if not isinstance(scopes, list):
        return False
    pair = {left, right}
    for scope in scopes:
        if not isinstance(scope, Mapping) or scope.get("page") != page:
            continue
        if scope.get("kind") not in (None, kind):
            continue
        elements = scope.get("elements")
        if isinstance(elements, list) and set(map(str, elements)) == pair:
            return True
    return False


def _overlap_issue(page: int, kind: str, left: tuple[str, str, tuple[float, float, float, float]], right: tuple[str, str, tuple[float, float, float, float]], intersection: tuple[float, float, float, float]) -> OfficeIssue:
    return OfficeIssue(
        "visual_overlap",
        "two text-bearing elements overlap",
        page=page,
        kind=kind,
        elements=(left[0], right[0]),
        texts=(left[1], right[1]),
        bboxes=(left[2], right[2]),
        intersection=intersection,
    )


def _round_box(values: Any) -> tuple[float, float, float, float]:
    return tuple(round(float(value), 3) for value in values)  # type: ignore[return-value]


def _cell_label(row: int, column: int) -> str:
    name = ""
    number = column + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        name = chr(65 + remainder) + name
    return f"{name}{row + 1}"
